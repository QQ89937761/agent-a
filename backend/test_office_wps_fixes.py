"""Real OOXML regression tests: no Office, COM, network or user applications."""
import asyncio
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from docx import Document
from PIL import Image

from backend.core.tools import ToolRegistry
from backend.core.text_normalize import normalize_cn_punctuation
from backend.core.agent import _remote_confirmation_categories, _looks_irreversible
from backend.modules.office import office_module, docgen_module, fs_module
from backend.modules.office.office_files import atomic_office_save, validate_plain_ooxml
from backend.modules.office.office_reading import read_excel_file


class OfficeWpsFixTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.registry = ToolRegistry()
        with patch.object(office_module, "_com_available", return_value=False):
            office_module.register_tools(self.registry)
        docgen_module.register_tools(self.registry)
        fs_module.register_tools(self.registry)
        self.backups = patch("backend.core.isolation_workspace.root", return_value=self.root / "isolation")
        self.backups.start()
        self.addCleanup(self.backups.stop)

    async def test_pdf_tools_read_pages_and_report_missing_file(self):
        from pypdf import PdfWriter
        path = self.root / "manual.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=400)
        with path.open("wb") as stream:
            writer.write(stream)
        self.assertEqual({"path": str(path), "pages": 1}, await self.registry.call("pdf_list_pages", {"path": str(path)}))
        text = await self.registry.call("pdf_read_text", {"path": str(path)})
        self.assertIn("第 1 页", text)
        with self.assertRaisesRegex(FileNotFoundError, "PDF 文件不存在"):
            await self.registry.call("pdf_list_pages", {"path": str(self.root / "missing.pdf")})

    async def test_word_batch_replacements_save_full_document_and_reject_missing(self):
        path = self.root / "draft.docx"
        doc = Document()
        doc.add_paragraph("原文第一段")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "原文表格"
        doc.save(path)
        result = await self.registry.call("word_apply_replacements", {"path": str(path), "replacements": [
            {"find": "原文", "replace": "修订"},
            {"find": "第一段", "replace": "全文"},
        ]})
        self.assertIn("已直接修改原文档并重新读取核验", result)
        edited = Document(path)
        self.assertEqual("修订全文", edited.paragraphs[0].text)
        self.assertEqual("修订表格", edited.tables[0].cell(0, 0).text)
        with self.assertRaisesRegex(FileNotFoundError, "Word 文件不存在"):
            await self.registry.call("word_apply_replacements", {"path": str(self.root / "missing.docx"), "replacements": [{"find": "a", "replace": "b"}]})

    async def test_word_batch_replacements_with_no_matches_are_nonfatal_and_do_not_save(self):
        path = self.root / "no-match.docx"
        doc = Document()
        doc.add_paragraph("第七章 原始正文")
        doc.save(path)
        before = path.read_bytes()
        result = await self.registry.call("word_apply_replacements", {
            "path": str(path), "replacements": [{"find": "不存在的锚点", "replace": "新内容"}],
        })
        self.assertIn("未找到任何待替换文字", result)
        self.assertIn("原 Word 文档未修改、未保存", result)
        self.assertEqual(before, path.read_bytes())
        self.assertFalse((self.root / "isolation/backups/office").exists())

    async def test_word_paragraph_listing_and_anchor_insertion(self):
        path = self.root / "anchored.docx"
        doc = Document()
        doc.add_paragraph("第七章 施工准备", style="Heading 1")
        doc.add_paragraph("原有内容")
        doc.add_paragraph("第八章 应急管理", style="Heading 1")
        doc.save(path)
        index = await self.registry.call("word_list_paragraphs", {"path": str(path)})
        self.assertEqual(3, index["total"])
        self.assertEqual("第八章 应急管理", index["paragraphs"][2]["text"])
        result = await self.registry.call("word_insert_paragraphs", {
            "path": str(path), "anchor_before": "第八章 应急管理", "paragraphs": ["补充的第七章与第八章之间内容"],
        })
        self.assertIn("已向原 Word 文档第 2 段位置插入 1 段", result)
        self.assertEqual(["第七章 施工准备", "原有内容", "补充的第七章与第八章之间内容", "第八章 应急管理"],
                         [paragraph.text for paragraph in Document(path).paragraphs])

    async def test_word_anchor_after_and_ambiguous_anchor_do_not_corrupt_document(self):
        path = self.root / "anchor-after.docx"
        doc = Document()
        doc.add_paragraph("章节标题")
        doc.add_paragraph("重复标题")
        doc.add_paragraph("重复标题")
        doc.save(path)
        result = await self.registry.call("word_insert_paragraphs", {
            "path": str(path), "anchor_after": "章节标题", "paragraphs": ["插在标题后"]
        })
        self.assertIn("第 1 段位置插入", result)
        self.assertEqual(["章节标题", "插在标题后", "重复标题", "重复标题"],
                         [paragraph.text for paragraph in Document(path).paragraphs])

        before = path.read_bytes()
        result = await self.registry.call("word_insert_paragraphs", {
            "path": str(path), "anchor_before": "重复标题", "paragraphs": ["不得插入"]
        })
        self.assertIn("标题锚点匹配到 2 个正文段落", result)
        self.assertIn("未修改、未保存", result)
        self.assertEqual(before, path.read_bytes())

    async def test_word_append_paragraphs_writes_to_original_and_creates_backup(self):
        path = self.root / "append.docx"
        doc = Document()
        doc.add_paragraph("原始正文")
        doc.save(path)
        before = path.read_bytes()
        result = await self.registry.call("word_append_paragraphs", {
            "path": str(path), "paragraphs": ["新增第一段", "新增第二段"],
        })
        self.assertIn("已向原 Word 文档追加 2 段并重新读取核验", result)
        self.assertEqual(["原始正文", "新增第一段", "新增第二段"], [p.text for p in Document(path).paragraphs])
        backups = list((self.root / "isolation/backups/office").glob("*.docx"))
        self.assertEqual(1, len(backups))
        self.assertEqual(before, backups[0].read_bytes())
        with self.assertRaisesRegex(FileNotFoundError, "Word 文件不存在"):
            await self.registry.call("word_append_paragraphs", {
                "path": str(self.root / "missing.docx"), "paragraphs": ["不会写入"],
            })

    async def test_pdf_extract_images_creates_requested_directory(self):
        from pypdf import PdfWriter
        path = self.root / "empty.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=400)
        with path.open("wb") as stream:
            writer.write(stream)
        output = self.root / "pdf-images"
        self.assertEqual([], await self.registry.call("pdf_extract_images", {
            "path": str(path), "output_dir": str(output),
        }))
        self.assertTrue(output.is_dir())

    async def test_pdf_render_pages_outputs_png_without_manual_export(self):
        from pypdf import PdfWriter
        path = self.root / "render.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=400)
        with path.open("wb") as stream:
            writer.write(stream)
        output = self.root / "rendered"
        result = await self.registry.call("pdf_render_pages", {
            "path": str(path), "output_dir": str(output), "pages": [1], "dpi": 72,
        })
        self.assertEqual([str(output / "page-001.png")], result)
        self.assertTrue((output / "page-001.png").read_bytes().startswith(b"\x89PNG"))

    async def test_word_insert_paragraphs_updates_original_and_verifies(self):
        path = self.root / "insert.docx"
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("第三段")
        doc.save(path)
        result = await self.registry.call("word_insert_paragraphs", {
            "path": str(path), "index": 1, "paragraphs": ["第二段", "第二段补充"],
        })
        self.assertIn("已向原 Word 文档第 1 段位置插入 2 段并重新读取核验", result)
        self.assertEqual(["第一段", "第二段", "第二段补充", "第三段"], [p.text for p in Document(path).paragraphs])

    async def test_word_append_blocks_writes_structured_content_to_original(self):
        path = self.root / "blocks.docx"
        doc = Document()
        doc.add_paragraph("原始正文")
        doc.save(path)
        result = await self.registry.call("word_append_blocks", {
            "path": str(path),
            "blocks": [
                {"type": "chapter", "text": "补充章节"},
                {"type": "body", "text": "补充正文"},
                {"type": "table", "caption": "表1 补充表", "headers": ["项目", "结果"], "rows": [["A", "已完成"]]},
            ],
        })
        self.assertIn("已向原 Word 文档追加 3 个结构块并重新读取核验", result)
        edited = Document(path)
        text = "\n".join(paragraph.text for paragraph in edited.paragraphs)
        self.assertIn("补充章节", text)
        self.assertIn("补充正文", text)
        self.assertEqual("已完成", edited.tables[-1].cell(1, 1).text)

    async def test_xlsx_edit_append_backup_and_clear_without_com(self):
        path = self.root / "book.xlsx"
        wb = Workbook()
        wb.active.title = "数据"
        wb.active["A1"] = "old"
        wb.save(path)
        before = path.read_bytes()
        await self.registry.call("excel_write_cells", {"path": str(path), "sheet": " 数据 ",
                                  "cells": {"A1:B2": [["new", 2], [3, 4]]}})
        backups = list((self.root / "isolation/backups/office").glob("*.xlsx"))
        self.assertEqual(1, len(backups))
        self.assertEqual(before, backups[0].read_bytes())
        self.assertEqual((("new", 2), (3, 4)), read_excel_file(path))
        await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"B2": None}})
        self.assertIsNone(read_excel_file(path, cell_range="B2"))
        result = await self.registry.call("excel_append_row", {"path": str(path), "values": ["last"]})
        self.assertIn("第 3 行", result)
        self.assertEqual("last", read_excel_file(path, cell_range="A3"))

    @staticmethod
    def rewrite_zip(path, changes):
        with ZipFile(path) as archive:
            parts = [(info, changes.get(info.filename, archive.read(info.filename))) for info in archive.infolist()]
        with ZipFile(path, "w") as archive:
            for info, content in parts:
                archive.writestr(info, content)

    async def test_filter_print_names_and_shared_formulas_round_trip(self):
        from openpyxl.workbook.defined_name import DefinedName
        path = self.root / "shared.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工程数据"
        ws.append([1, 2, "=A1+B1"])
        ws.append([3, 4, "=A2+B2"])
        ws.auto_filter.ref = "A1:C2"
        ws.print_area = "A1:C2"
        ws.print_title_rows = "1:1"
        wb.defined_names.add(DefinedName("本地范围", attr_text="'工程数据'!$A$1:$B$2"))
        wb.save(path)
        with ZipFile(path) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode().replace(
                "<f>A1+B1</f>", '<f t="shared" ref="C1:C2" si="0">A1+B1</f>').replace(
                "<f>A2+B2</f>", '<f t="shared" si="0"/>')
        self.rewrite_zip(path, {"xl/worksheets/sheet1.xml": xml})
        validate_plain_ooxml(path)
        await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"A1": 10}})
        await self.registry.call("excel_append_row", {"path": str(path), "values": [5, 6]})
        wb = load_workbook(path)
        try:
            ws = wb.active
            self.assertEqual((10, 5), (ws["A1"].value, ws["A3"].value))
            self.assertEqual(("=A1+B1", "=A2+B2"), (ws["C1"].value, ws["C2"].value))
            self.assertEqual("A1:C2", ws.auto_filter.ref)
            self.assertIn("$A$1:$C$2", str(ws.print_area))
            self.assertEqual("$1:$1", ws.print_title_rows)
            self.assertIn("本地范围", wb.defined_names)
        finally:
            wb.close()
        # Editing a shared-formula member must leave the other formula intact.
        await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"C1": 7}})
        wb = load_workbook(path)
        try:
            self.assertEqual("=A2+B2", wb.active["C2"].value)
        finally:
            wb.close()

    async def test_whole_column_print_area_and_reserved_names_preserved_exactly(self):
        from lxml import etree as ET
        path = self.root / "whole-column-print.xlsx"
        wb = Workbook()
        wb.active["A1"] = "old"
        wb.active.print_area = "A1:C2"
        wb.save(path)
        with ZipFile(path) as archive:
            xml = archive.read("xl/workbook.xml").replace(b"$A$1:$C$2", b"$A:$X")
        self.rewrite_zip(path, {"xl/workbook.xml": xml})
        def names():
            with ZipFile(path) as archive:
                return ET.tostring(ET.fromstring(archive.read("xl/workbook.xml")).find(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedNames"))
        before = names()
        await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"A1": "new"}})
        self.assertEqual(before, names())

    async def test_risky_names_and_special_formulas_remain_blocked(self):
        from openpyxl.workbook.defined_name import DefinedName
        for formula in ['WEBSERVICE("https://example.test")', "'[external.xlsx]Sheet1'!$A$1"]:
            path = self.root / "risky-name.xlsx"
            wb = Workbook()
            wb.defined_names.add(DefinedName("Risky", attr_text=formula))
            wb.save(path)
            before = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "外部或未知公式"):
                await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"B1": 2}})
            self.assertEqual(before, path.read_bytes())
        for kind in ("array", "dataTable", "unknown"):
            path = self.root / "special.xlsx"
            wb = Workbook()
            wb.active["A1"] = "=SUM(B1:B2)"
            wb.save(path)
            with ZipFile(path) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml").replace(b"<f>", f'<f t="{kind}">'.encode())
            self.rewrite_zip(path, {"xl/worksheets/sheet1.xml": xml})
            with self.assertRaisesRegex(ValueError, "特殊公式"):
                validate_plain_ooxml(path)

    async def test_risky_shared_formula_and_active_word_parts_still_block(self):
        path = self.root / "risky-shared.xlsx"
        wb = Workbook()
        wb.active["A1"] = '=WEBSERVICE("https://example.test")'
        wb.save(path)
        with ZipFile(path) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").replace(b"<f>", b'<f t="shared" si="0" ref="A1:A2">')
        self.rewrite_zip(path, {"xl/worksheets/sheet1.xml": xml})
        with self.assertRaisesRegex(ValueError, "外部或未知公式"):
            validate_plain_ooxml(path)
        for name in ("word/vbaProject.bin", "word/activeX/activeX1.xml", "_xmlsignatures/sig1.xml"):
            path = self.root / "active.docx"
            Document().save(path)
            with ZipFile(path, "a") as archive:
                archive.writestr(name, b"active")
            before = path.read_bytes()
            with self.assertRaises(ValueError):
                await self.registry.call("word_find_replace", {"path": str(path), "find": "old", "replace": "new"})
            self.assertEqual(before, path.read_bytes())

    async def test_word_objects_and_untouched_package_parts_are_preserved(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        path = self.root / "objects.docx"
        image = self.root / "figure.png"
        Image.new("RGB", (4, 4), "white").save(image)
        doc = Document()
        p = doc.add_paragraph("old")
        p.add_run().add_tab()
        p.add_run("old").bold = True
        p.add_run().add_break()
        p.add_run("old")
        p.add_run().add_picture(str(image))
        p.add_run("old")
        for name, content_type, rel, payload in [
            ("/word/embeddings/oleObject1.bin", "application/vnd.openxmlformats-officedocument.oleObject", "oleObject", b"untouched OLE payload"),
            ("/customXml/item99.xml", "application/xml", "customXml", b'<wpsMetadata xmlns="urn:kso">keep</wpsMetadata>'),
        ]:
            part = Part(PackURI(name), content_type, payload, doc.part.package)
            doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + rel)
        doc.save(path)
        with ZipFile(path, "a") as archive:
            archive.writestr("customXml/orphan-metadata.xml", b"<unreferenced>keep</unreferenced>")
        with ZipFile(path) as archive:
            before_parts = {n: archive.read(n) for n in archive.namelist() if n.startswith(("word/media/", "word/embeddings/", "customXml/"))}
        before = path.read_bytes()
        result = await self.registry.call("word_find_replace", {"path": str(path), "find": "old", "replace": "new"})
        self.assertIn("4 处", result)
        doc = Document(path)
        self.assertEqual("new\tnew\nnewnew", doc.paragraphs[0].text)
        self.assertEqual(1, len(doc.inline_shapes))
        self.assertTrue(doc.paragraphs[0].runs[2].bold)
        with ZipFile(path) as archive:
            for name, payload in before_parts.items():
                self.assertEqual(payload, archive.read(name), name)
        self.assertEqual(before, next((self.root / "isolation/backups/office").glob("*.docx")).read_bytes())

    async def test_word_no_match_does_not_save_or_backup(self):
        path = self.root / "no-match.docx"
        doc = Document()
        doc.add_paragraph("unchanged")
        doc.save(path)
        before, mtime = path.read_bytes(), path.stat().st_mtime_ns
        with patch("backend.modules.office.office_files.atomic_office_save") as save:
            result = await self.registry.call("word_find_replace", {"path": str(path), "find": "absent", "replace": "new"})
        save.assert_not_called()
        self.assertIn("未在文档中找到", result)
        self.assertEqual((before, mtime), (path.read_bytes(), path.stat().st_mtime_ns))
        self.assertFalse((self.root / "isolation/backups/office").exists())

    async def test_word_does_not_match_across_tabs_breaks_or_pictures(self):
        from docx.oxml import OxmlElement
        image = self.root / "boundary.png"
        Image.new("RGB", (4, 4), "white").save(image)
        for kind in ("tab", "br", "drawing", "pict"):
            path = self.root / f"boundary-{kind}.docx"
            doc = Document()
            p = doc.add_paragraph("old")
            run = p.add_run()
            if kind == "drawing":
                run.add_picture(str(image))
            else:
                run._r.append(OxmlElement("w:" + kind))
            p.add_run("text")
            doc.save(path)
            before = path.read_bytes()
            result = await self.registry.call("word_find_replace", {"path": str(path), "find": "oldtext", "replace": "wrong"})
            self.assertIn("未在文档中找到", result)
            self.assertEqual(before, path.read_bytes())
            result = await self.registry.call("word_find_replace", {"path": str(path), "find": "old", "replace": "new"})
            self.assertIn("1 处", result)
            self.assertEqual(1, len(list(Document(path).paragraphs[0]._p.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + kind))))

    async def test_word_fields_still_block_and_leave_original_untouched(self):
        from docx.oxml import OxmlElement
        path = self.root / "field.docx"
        doc = Document()
        doc.add_paragraph("old")
        p = doc.add_paragraph("old")
        p.add_run()._r.append(OxmlElement("w:fldChar"))
        doc.save(path)
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "域代码"):
            await self.registry.call("word_find_replace", {"path": str(path), "find": "old", "replace": "new"})
        self.assertEqual(before, path.read_bytes())
        self.assertFalse((self.root / "isolation/backups/office").exists())

    async def test_append_empty_styled_sheet_starts_at_one(self):
        path = self.root / "empty.xlsx"
        wb = Workbook()
        wb.active["A500"].number_format = "0.00"
        wb.save(path)
        result = await self.registry.call("excel_append_row", {"path": str(path), "values": ["first"]})
        self.assertIn("第 1 行", result)
        self.assertEqual("first", read_excel_file(path))

    async def test_list_sheets_trim_and_missing_name(self):
        path = self.root / "sheet.xlsx"
        wb = Workbook()
        wb.active.title = "数据"
        wb.active["A1"] = "value"
        wb.active["F50"].number_format = "0.00"
        wb.create_sheet("汇总")
        wb.save(path)
        result = await self.registry.call("excel_list_sheets", {"path": str(path)})
        self.assertEqual(["数据", "汇总"], [item["name"] for item in result["sheets"]])
        self.assertEqual("value", read_excel_file(path, " 数据 "))
        self.assertEqual((("value", None), (None, None)), read_excel_file(path, cell_range="A1:B2"))
        with self.assertRaisesRegex(ValueError, "数据.*汇总"):
            read_excel_file(path, "wrong")

    async def test_unsafe_existing_workbook_does_not_fall_back_to_com(self):
        path = self.root / "external.xlsx"
        wb = Workbook()
        wb.active["A1"] = '=WEBSERVICE("https://example.test")'
        wb.save(path)
        before = path.read_bytes()
        with patch.object(office_module, "_dispatch_office") as dispatch:
            with self.assertRaisesRegex(ValueError, "外部或未知公式"):
                await self.registry.call("excel_write_cells", {"path": str(path), "cells": {"B1": 2}})
        dispatch.assert_not_called()
        self.assertEqual(before, path.read_bytes())

    async def test_word_cross_run_and_table_replace_preserves_unaffected_style(self):
        path = self.root / "word.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("keep ").italic = True
        paragraph.add_run("ol").bold = True
        paragraph.add_run("d tail")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "old table"
        doc.save(path)
        before = path.read_bytes()
        await self.registry.call("word_find_replace", {"path": str(path), "find": "old", "replace": "new"})
        result = Document(path)
        self.assertEqual("keep new tail", result.paragraphs[0].text)
        self.assertTrue(result.paragraphs[0].runs[0].italic)
        self.assertEqual("new table", result.tables[0].cell(0, 0).text)
        self.assertEqual(before, next((self.root / "isolation/backups/office").glob("*.docx")).read_bytes())

    async def test_generators_multisheet_image_and_explicit_overwrite(self):
        excel = self.root / "multi.xlsx"
        args = {"path": str(excel), "sheets": [{"name": "数据", "headers": ["名称", "数量"], "rows": [["钢筋", 2]]},
                                                {"name": "汇总", "rows": [["=SUM(1,2)"]]}]}
        await self.registry.call("create_excel_file", args)
        original = excel.read_bytes()
        with self.assertRaisesRegex(ValueError, "避免覆盖"):
            await self.registry.call("create_excel_file", args)
        await self.registry.call("create_excel_file", {**args, "overwrite": True})
        self.assertEqual(original, next((self.root / "isolation/backups/office").glob("*.xlsx")).read_bytes())
        wb = load_workbook(excel)
        self.assertEqual(["数据", "汇总"], wb.sheetnames)
        self.assertEqual("=SUM(1,2)", wb["汇总"]["A1"].value)
        self.assertGreaterEqual(wb["数据"].column_dimensions["A"].width, 10)
        wb.close()
        image = self.root / "figure.png"
        Image.new("RGB", (1200, 800), "white").save(image)
        word = self.root / "figure.docx"
        await self.registry.call("create_word_doc", {"path": str(word), "blocks": [
            {"type": "figure", "image_path": str(image), "text": "图1 示例"}]})
        doc = Document(word)
        self.assertEqual(1, len(doc.inline_shapes))
        self.assertIn("图1 示例", "".join(p.text for p in doc.paragraphs))
        await self.registry.call("create_word_doc", {"path": str(word), "paragraphs": ["new"], "overwrite": True})
        self.assertEqual(1, len(list((self.root / "isolation/backups/office").glob("*.docx"))))

    async def test_invalid_multisheet_and_missing_image_leave_no_output(self):
        for sheets in ([], [{"name": "Bad/Name"}], [{"name": "A"}, {"name": "a"}], [{"name": "A", "rows": [1]}]):
            path = self.root / "invalid.xlsx"
            with self.assertRaises(ValueError):
                await self.registry.call("create_excel_file", {"path": str(path), "sheets": sheets})
            self.assertFalse(path.exists())
        path = self.root / "invalid.docx"
        with self.assertRaisesRegex(ValueError, "图片不存在"):
            await self.registry.call("create_word_doc", {"path": str(path), "blocks": [{"type": "figure", "path": str(self.root / "missing.png")}]})
        self.assertFalse(path.exists())

    async def test_binary_text_read_has_actionable_guide(self):
        for suffix, tool in ((".xlsx", "excel_read_all"), (".docx", "word_read_text"), (".pptx", "ppt_read_text")):
            path = self.root / ("binary" + suffix)
            path.write_bytes(b"PK binary")
            with self.assertRaisesRegex(ValueError, tool):
                await self.registry.call("fs_read_text", {"path": str(path)})

    def test_atomic_save_failure_and_no_clobber_preserve_original(self):
        path = self.root / "original.xlsx"
        path.write_bytes(b"original")
        def fail(temporary):
            Path(temporary).write_bytes(b"partial")
            raise OSError("disk full")
        with self.assertRaisesRegex(OSError, "disk full"):
            atomic_office_save(path, fail, overwrite=True)
        self.assertEqual(b"original", path.read_bytes())
        self.assertFalse(list(self.root.glob(".xiaoa-office-*")))
        new = self.root / "race.xlsx"
        def race(temporary):
            Path(temporary).write_bytes(b"ours")
            new.write_bytes(b"theirs")
        with self.assertRaises(ValueError):
            atomic_office_save(new, race)
        self.assertEqual(b"theirs", new.read_bytes())

    def test_com_wps_progids_and_bounded_busy_retry(self):
        path = self.root / "active.docx"
        document = Mock(FullName=str(path))
        documents = Mock(Count=1)
        documents.Item.return_value = document
        app = Mock(Documents=documents)
        client = Mock()
        client.GetActiveObject.side_effect = lambda progid: app if progid == "KWPS.Application" else (_ for _ in ()).throw(RuntimeError("missing"))
        self.assertEqual((app, document, False), office_module._find_open_word_document(client, path))
        class Busy(Exception):
            hresult = -2147418111
        call = Mock(side_effect=[Busy(), Busy(), "ok"])
        with patch.object(office_module.time, "sleep") as sleep:
            self.assertEqual("ok", office_module._retry_com_busy(call))
            self.assertEqual(2, sleep.call_count)
        with patch.object(office_module.time, "sleep"):
            with self.assertRaises(Busy):
                office_module._retry_com_busy(Mock(side_effect=Busy()))

    def test_technical_spans_and_multisheet_risk(self):
        for technical in (r"D:\办公\结果.docx", "https://example.com/a?q=1,b=2", "SUM(A1:A10)", "sin(x)"):
            self.assertIn(technical, normalize_cn_punctuation("保存到" + technical))
        self.assertEqual("说明：完成，继续。", normalize_cn_punctuation("说明:完成,继续。"))
        self.assertEqual({"overwrite"}, _remote_confirmation_categories("excel_write_cells", {"path": "a.xlsx", "cells": {"A1": "plain"}}))
        self.assertEqual({"overwrite", "external_code"}, _remote_confirmation_categories("excel_write_cells", {"path": "a.xlsx", "cells": {"A1:B1": [[1, '=WEBSERVICE("https://test")']]}}))
        self.assertEqual({"overwrite", "external_code"}, _remote_confirmation_categories("word_find_replace", {"path": "a.docm"}))
        self.assertEqual({"overwrite"}, _remote_confirmation_categories("create_word_doc", {"overwrite": True}))
        args = {"sheets": [{"name": "safe", "rows": [[1]]}, {"name": "unsafe", "rows": [['=EXEC("format D:")']]}]}
        self.assertEqual({"external_code"}, _remote_confirmation_categories("create_excel_file", args))
        self.assertTrue(_looks_irreversible("create_excel_file", args))

    async def test_ppt_respects_slide_order_reads_tables_and_notes(self):
        path = self.root / "slides.pptx"
        rel = "http://schemas.openxmlformats.org/package/2006/relationships"
        office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        pres = "http://schemas.openxmlformats.org/presentationml/2006/main"
        draw = "http://schemas.openxmlformats.org/drawingml/2006/main"
        with ZipFile(path, "w") as archive:
            archive.writestr("ppt/presentation.xml", f'<p:presentation xmlns:p="{pres}" xmlns:r="{office}"><p:sldIdLst><p:sldId id="257" r:id="second"/><p:sldId id="256" r:id="first"/></p:sldIdLst></p:presentation>')
            archive.writestr("ppt/_rels/presentation.xml.rels", f'<Relationships xmlns="{rel}"><Relationship Id="first" Type="{office}/slide" Target="slides/slide1.xml"/><Relationship Id="second" Type="{office}/slide" Target="slides/slide2.xml"/></Relationships>')
            for number, text in ((1, "第一页文件，排在第二"), (2, "第二页文件，排在第一")):
                archive.writestr(f"ppt/slides/slide{number}.xml", f'<p:sld xmlns:p="{pres}" xmlns:a="{draw}"><a:p><a:r><a:t>{text}</a:t></a:r></a:p><a:tbl><a:tr><a:tc><a:p><a:r><a:t>表格单元格</a:t></a:r></a:p></a:tc></a:tr></a:tbl></p:sld>')
            archive.writestr("ppt/slides/_rels/slide2.xml.rels", f'<Relationships xmlns="{rel}"><Relationship Id="note" Type="{office}/notesSlide" Target="../notesSlides/notesSlide1.xml"/></Relationships>')
            archive.writestr("ppt/notesSlides/notesSlide1.xml", f'<p:notes xmlns:p="{pres}" xmlns:a="{draw}"><p:sp><a:p><a:r><a:t>讲者备注</a:t></a:r></a:p></p:sp><p:sp><p:ph type="sldNum"/><a:p><a:r><a:t>999</a:t></a:r></a:p></p:sp></p:notes>')
        result = await self.registry.call("ppt_read_text", {"path": str(path)})
        self.assertLess(result.index("第二页文件"), result.index("第一页文件"))
        self.assertIn("表格单元格", result)
        self.assertIn("讲者备注", result)
        self.assertIn("【第 2 页】", result)
        self.assertNotIn("999", result)
        self.assertEqual(set(), _remote_confirmation_categories("ppt_read_text", {"path": str(path)}))

    async def test_lock_marker_blocks_generators_and_preserves_destination(self):
        for tool, suffix in (("create_word_doc", ".docx"), ("create_excel_file", ".xlsx")):
            path = self.root / ("locked" + suffix)
            path.write_bytes(b"original")
            path.with_name("~$" + path.name).touch()
            with self.assertRaisesRegex(RuntimeError, "占用标记"):
                await self.registry.call(tool, {"path": str(path), "overwrite": True})
            self.assertEqual(b"original", path.read_bytes())
        self.assertFalse((self.root / "isolation/backups/office").exists())

    def test_backup_failure_never_changes_original(self):
        path = self.root / "book.xlsx"
        path.write_bytes(b"original")
        with patch("backend.modules.office.office_files.shutil.copy2", side_effect=OSError("backup failed")):
            with self.assertRaisesRegex(OSError, "backup failed"):
                atomic_office_save(path, lambda temporary: Path(temporary).write_bytes(b"new"), overwrite=True)
        self.assertEqual(b"original", path.read_bytes())
        self.assertFalse(list(self.root.glob(".xiaoa-office-*")))

    def test_isolation_cleanup_preserves_office_backups(self):
        from backend.core import isolation_workspace
        folder = isolation_workspace.prepare()
        backup = folder / "backups/office/original.xlsx"
        backup.parent.mkdir(parents=True)
        backup.write_bytes(b"backup")
        cache = folder / "tmp/cache.bin"
        cache.write_bytes(b"cache")
        with patch.object(isolation_workspace, "usage_bytes", return_value=129 * 1024 * 1024):
            isolation_workspace.cleanup(128)
        self.assertEqual(b"backup", backup.read_bytes())
        self.assertFalse(cache.exists())
