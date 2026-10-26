"""文档生成模块(不依赖微软 Office,WPS 也能打开)。

用 python-docx / openpyxl 直接生成真正的 .docx / .xlsx 文件,默认不启动 Office/WPS。
这样无论用户装的是微软 Office 还是 WPS,"新建 Word 文档""新建 Excel 表格"都能用。

跨平台可导入。
"""
from __future__ import annotations

import uuid
from pathlib import Path

from backend.core.text_normalize import normalize_cn_punctuation
from backend.modules.office.excel_formulas import is_formula
from backend.modules.office.office_files import atomic_office_save
from backend.core.tools import ToolRegistry
from backend.core import isolation_workspace
from backend.modules.computer.code_module import _reject_protected


def _default_docs_dir() -> Path:
    return isolation_workspace.workspace()


def _resolve_path(path: str | None, default_name: str) -> Path:
    if not path or not path.strip():
        p = _default_docs_dir() / default_name
        if p.exists():
            p = p.with_name(f"{p.stem}-{uuid.uuid4().hex[:12]}{p.suffix}")
        return p
    p = Path(path)
    if not p.is_absolute():
        p = _default_docs_dir() / p
    p = p.resolve()
    _reject_protected(p)
    return p



def register_tools(registry: ToolRegistry) -> None:
    @registry.tool(
        name="create_word_doc",
        description=(
            "新建并保存一个 Word 文档(.docx)，不启动 WPS/Word。"
            "**默认自动套用用户的 Word 排版规范**(A4/四周 25mm 页边距/正文小四宋体/固定行距 22 磅/"
            "首行缩进 2 字符/章标题小三居中加粗/表头浅蓝底黑字/页脚页码),不需要用户每次交代;"
            "只有用户明确要求别的版式时才另说。"
            "写正式文档请用 blocks 按结构给内容(章/节/正文/条款/表格/图题),比只给 paragraphs 更合规。"
            "默认拒绝覆盖；显式 overwrite=true 时先备份再原子覆盖。"
        ),
        parameters={
            "type": "object",
            "properties": {
                "overwrite": {"type": "boolean", "default": False, "description": "显式允许覆盖，先备份原文件；默认拒绝覆盖"},
                "path": {"type": "string", "description": "保存路径,可省略(默认放到隔离工作区，不要猜测用户文档/桌面路径)。也可只给文件名"},
                "title": {"type": "string", "description": "文档标题(按章标题排版),可省略"},
                "paragraphs": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "正文段落列表(按正文排版),可省略(空文档)",
                },
                "blocks": {
                    "type": "array",
                    "items": {"type": "object"},
                    "description": (
                        "结构化内容(推荐,优先于 title/paragraphs)。每项:"
                        "{type:'chapter'|'section'|'body'|'clause', text:'…'} 或 "
                        "{type:'table', caption:'表1 …', headers:[…], rows:[[…]]} 或 "
                        "{type:'figure', image_path:'图片绝对路径', text:'图1 …'} 或 {type:'figure_caption', text:'图1 …'}"
                    ),
                },
                "header": {"type": "string", "description": "页眉文字(一般填项目名称,左对齐);省略则不写页眉"},
            },
            "required": [],
        },
        side_effect=True,
    )
    async def create_word_doc(
        path: str | None = None,
        title: str | None = None,
        paragraphs: list[str] | None = None,
        blocks: list[dict] | None = None,
        header: str | None = None,
        overwrite: bool = False,
    ) -> str:
        from docx import Document

        from backend.modules.office import word_style

        p = _resolve_path(path, "新建文档.docx")
        if p.suffix.lower() != ".docx":
            p = p.with_suffix(".docx")
        if not isinstance(overwrite, bool):
            raise ValueError("overwrite 必须是布尔值")
        if p.exists() and not overwrite:
            raise ValueError(f"目标文件已存在，为避免覆盖未写入: {p}")
        p.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        # 默认套用长期记忆里的 Word 排版规范(字体/字号/行距/页边距/页眉页脚)
        word_style.apply_default_style(doc)
        if header:
            word_style.set_header(doc, normalize_cn_punctuation(str(header)))
        word_style.add_page_number_footer(doc)

        def clean(value: object) -> str:
            # 模型写中文常混用半角标点,落盘前统一成中文全角
            return normalize_cn_punctuation("" if value is None else str(value))

        def table_array(block: dict, name: str) -> list:
            value = block.get(name)
            if value is None:
                return []
            if not isinstance(value, list):
                raise ValueError(f"Word 表格字段 {name} 必须是数组")
            return value

        if blocks:
            for block in blocks:
                if not isinstance(block, dict):
                    continue
                kind = str(block.get("type") or "body").lower()
                text = clean(block.get("text"))
                if kind == "chapter":
                    word_style.add_chapter(doc, text)
                elif kind == "section":
                    word_style.add_section(doc, text)
                elif kind == "clause":
                    word_style.add_clause(doc, text)
                elif kind == "table":
                    headers = table_array(block, "headers")
                    rows = table_array(block, "rows")
                    if any(not isinstance(row, list) for row in rows):
                        raise ValueError("Word 表格字段 rows 必须是二维数组")
                    word_style.add_table(
                        doc,
                        [clean(h) for h in headers],
                        [[clean(c) for c in row] for row in rows],
                        caption=clean(block.get("caption")),
                    )
                elif kind in ("figure", "figure_caption"):
                    image = block.get("image_path") or block.get("path")
                    if kind == "figure" and image:
                        from docx.shared import Cm
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        from docx.image.image import Image
                        from backend.modules.office.fs_module import _resolve
                        image_path = _resolve(str(image))
                        if not image_path.is_file():
                            raise ValueError(f"图片不存在: {image_path}")
                        info = Image.from_file(str(image_path))
                        section = doc.sections[-1]
                        max_width = min(Cm(14.5), section.page_width - section.left_margin - section.right_margin)
                        max_height = min(Cm(18), section.page_height - section.top_margin - section.bottom_margin)
                        scale = min(1, max_width / info.width, max_height / info.height)
                        paragraph = doc.add_paragraph()
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.add_run().add_picture(str(image_path), width=int(info.width * scale), height=int(info.height * scale))
                    word_style.add_figure_caption(doc, text or clean(block.get("caption")))
                else:
                    word_style.add_body(doc, text)
        else:
            if title:
                word_style.add_chapter(doc, clean(title))
            for para in paragraphs or []:
                word_style.add_body(doc, clean(para))
        from backend.modules.office.office_module import _check_closed_for_save
        backup = atomic_office_save(p, doc.save, overwrite=overwrite,
                                   check=lambda: _check_closed_for_save(p, word=True))
        return f"已新建并保存 Word 文档(未启动 WPS/Word):{p}" + (f"；原文件备份: {backup}" if backup else "")

    @registry.tool(
        name="create_excel_file",
        description=(
            "新建并保存一个 Excel 表格(.xlsx)，不启动 WPS/Excel。"
            "可传入表头和若干行数据；普通数据和内置计算公式直接执行，外部连接或未知公式复用对话授权。"
            "默认拒绝覆盖；显式 overwrite=true 时先备份再原子覆盖。"
        ),
        parameters={
            "type": "object",
            "properties": {
                "overwrite": {"type": "boolean", "default": False, "description": "显式允许覆盖，先备份原文件；默认拒绝覆盖"},
                "path": {"type": "string", "description": "保存路径,可省略(默认放到隔离工作区，不要猜测用户文档/桌面路径)"},
                "sheets": {"type": "array", "items": {"type": "object"}, "description": "多工作表，每项含 name、headers、rows；不能与顶层 headers/rows 混用"},
                "headers": {"type": "array", "items": {"type": "string"}, "description": "表头,可省略"},
                "rows": {
                    "type": "array",
                    "items": {"type": "array", "items": {}},
                    "description": "数据行,每行是一个数组,可省略",
                },
            },
            "required": [],
        },
        side_effect=True,
    )
    async def create_excel_file(
        path: str | None = None, headers: list[str] | None = None, rows: list[list] | None = None,
        sheets: list[dict] | None = None, overwrite: bool = False,
    ) -> str:
        from openpyxl import Workbook

        if headers is not None and not isinstance(headers, list):
            raise ValueError("Excel 表头 headers 必须是数组")
        if rows is not None and (not isinstance(rows, list)
                                 or any(not isinstance(row, list) for row in rows)):
            raise ValueError("Excel 数据 rows 必须是二维数组")
        p = _resolve_path(path, "新建表格.xlsx")
        if p.suffix.lower() != ".xlsx":
            p = p.with_suffix(".xlsx")
        if not isinstance(overwrite, bool):
            raise ValueError("overwrite 必须是布尔值")
        if p.exists() and not overwrite:
            raise ValueError(f"目标文件已存在，为避免覆盖未写入: {p}")
        p.parent.mkdir(parents=True, exist_ok=True)

        def clean(value):
            # 公式是语法，不是中文正文；不能改写引号、参数分隔符或工作表引用。
            if isinstance(value, str) and not is_formula(value):
                return normalize_cn_punctuation(value)
            return value

        specifications = sheets if sheets is not None else [{"name": "Sheet", "headers": headers, "rows": rows}]
        if sheets is not None and (headers is not None or rows is not None):
            raise ValueError("sheets 不能与顶层 headers/rows 混用")
        if not isinstance(specifications, list) or not specifications:
            raise ValueError("sheets 必须是非空数组")
        names = set()
        import re
        for spec in specifications:
            if not isinstance(spec, dict):
                raise ValueError("每个工作表必须是对象")
            name = spec.get("name")
            if (not isinstance(name, str) or not name.strip() or len(name) > 31
                    or re.search(r"[\\/*?:\[\]]", name) or name.casefold() in names
                    or name.startswith("'") or name.endswith("'")):
                raise ValueError("工作表名无效或重复（不区分大小写，最多 31 字符）")
            names.add(name.casefold())
            if spec.get("headers") is not None and not isinstance(spec["headers"], list):
                raise ValueError("Excel 表头 headers 必须是数组")
            data = spec.get("rows")
            if data is not None and (not isinstance(data, list) or any(not isinstance(row, list) for row in data)):
                raise ValueError("Excel 数据 rows 必须是二维数组")
        wb = Workbook()
        wb.remove(wb.active)
        try:
            from openpyxl.utils import get_column_letter
            import unicodedata
            for spec in specifications:
                ws = wb.create_sheet(spec["name"])
                if spec.get("headers"):
                    ws.append([clean(value) for value in spec["headers"]])
                for row in spec.get("rows") or []:
                    ws.append([clean(value) for value in row])
                for column in ws.columns:
                    width = max((sum(2 if unicodedata.east_asian_width(char) in "WF" else 1
                                     for char in str(cell.value or "")) for cell in column), default=0)
                    ws.column_dimensions[get_column_letter(column[0].column)].width = min(60, max(10, width + 2))
            from backend.modules.office.office_module import _check_closed_for_save
            backup = atomic_office_save(p, wb.save, overwrite=overwrite,
                                       check=lambda: _check_closed_for_save(p))
        finally:
            wb.close()
        return f"已新建并保存 Excel 表格(未启动 WPS/Excel):{p}" + (f"；原文件备份: {backup}" if backup else "")
