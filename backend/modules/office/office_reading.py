"""不启动 Office 的只读文件解析；公式/宏/外链均不执行。"""
from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree


OFFICE_READ_TOOLS = {
    "excel_read_range", "excel_read_all", "excel_list_sheets", "word_read_text", "word_list_paragraphs", "ppt_read_text",
    "pdf_read_text", "pdf_list_pages",
}
_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
_WORD_SUFFIXES = {".docx", ".docm"}


def supports_file_read(name: str, args: dict[str, Any]) -> bool:
    path = args.get("path")
    if not isinstance(path, str):
        return False
    suffix = Path(path).suffix.lower()
    return (name in {"excel_read_range", "excel_read_all", "excel_list_sheets"} and suffix in _EXCEL_SUFFIXES
            or name == "word_read_text" and suffix in _WORD_SUFFIXES
            or name == "word_list_paragraphs" and suffix == ".docx"
            or name == "ppt_read_text" and suffix == ".pptx"
            or name in {"pdf_read_text", "pdf_list_pages"} and suffix == ".pdf")


def _pdf_reader(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("当前开发版缺少 PDF 解析组件 pypdf，请重新打包") from exc
    try:
        return PdfReader(str(path), strict=False)
    except Exception as exc:
        raise ValueError(f"PDF 无法读取或已损坏: {path} ({exc})") from exc


def read_pdf_file(path: Path) -> str:
    reader = _pdf_reader(path)
    pages = []
    for index, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = f"[第 {index} 页文字提取失败: {type(exc).__name__}]"
        pages.append(f"【第 {index} 页】\n{text.strip()}".rstrip())
    return "\n\n".join(pages)


def list_pdf_pages(path: Path) -> dict[str, object]:
    reader = _pdf_reader(path)
    return {"path": str(path), "pages": len(reader.pages)}


def extract_pdf_images(path: Path, output_dir: Path) -> list[str]:
    reader = _pdf_reader(path)
    output_dir.mkdir(parents=True, exist_ok=True)
    result = []
    for page_number, page in enumerate(reader.pages, 1):
        for image_number, image in enumerate(getattr(page, "images", ()), 1):
            target = output_dir / f"page-{page_number:03d}-image-{image_number:03d}{Path(image.name).suffix or '.bin'}"
            target.write_bytes(image.data)
            result.append(str(target))
    return result


def render_pdf_pages(path: Path, output_dir: Path, *, pages: list[int] | None = None, dpi: int = 150) -> list[str]:
    """Render PDF pages to PNGs so scanned pages can be inspected without manual export."""
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError("当前开发版缺少 PDF 页面渲染组件 PyMuPDF，请重新打包") from exc
    if not isinstance(dpi, int) or not 72 <= dpi <= 600:
        raise ValueError("dpi 必须是 72 到 600 之间的整数")
    if pages is not None and (not isinstance(pages, list) or any(not isinstance(page, int) or page < 1 for page in pages)):
        raise ValueError("pages 必须是从 1 开始的页码数组")
    output_dir.mkdir(parents=True, exist_ok=True)
    document = pymupdf.open(str(path))
    try:
        selected = pages if pages is not None else list(range(1, len(document) + 1))
        if any(page > len(document) for page in selected):
            raise ValueError(f"PDF 页码超出范围；当前共 {len(document)} 页")
        scale = dpi / 72
        matrix = pymupdf.Matrix(scale, scale)
        result = []
        for page_number in selected:
            image_path = output_dir / f"page-{page_number:03d}.png"
            document[page_number - 1].get_pixmap(matrix=matrix, alpha=False).save(str(image_path))
            result.append(str(image_path))
        return result
    finally:
        document.close()


def read_excel_file(path: Path, sheet: str | None = None, cell_range: str | None = None) -> Any:
    from openpyxl.utils.cell import range_boundaries

    bounds = None
    if cell_range is not None:
        bounds = range_boundaries(cell_range)
        if (any(value is None for value in bounds)
                or not 1 <= bounds[0] <= bounds[2] <= 16384
                or not 1 <= bounds[1] <= bounds[3] <= 1048576):
            raise ValueError("cell_range 必须是有效的有限单元格区域，如 A1:D20")
    # ponytail: 返回已保存的公式缓存；缺缓存为 None，不启动 Excel 重新计算。
    # 以字节流解析真实内容，不让链接目标的扩展名另行改变读取方式。
    from openpyxl import load_workbook
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            ws = select_sheet(workbook, sheet)
            if bounds is None:
                bounds = (ws.min_column or 1, ws.min_row or 1, ws.max_column or 1, ws.max_row or 1)
            min_col, min_row, max_col, max_row = bounds
            rows = tuple(ws.iter_rows(min_row=min_row, max_row=max_row,
                                      min_col=min_col, max_col=max_col, values_only=True))
            if cell_range is None:
                end = len(rows)
                while end and all(value is None for value in rows[end - 1]):
                    end -= 1
                rows = rows[:end]
                width = max((index + 1 for row in rows for index, value in enumerate(row)
                             if value is not None), default=0)
                rows = tuple(row[:width] for row in rows)
            if not rows:
                return None
            return rows[0][0] if len(rows) == 1 and len(rows[0]) == 1 else rows
        finally:
            workbook.close()


def read_word_file(path: Path) -> str:
    # 直接读主文档 XML，保留正文、表格、制表符和换行；不加载宏、模板或外部链接。
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    separators = {namespace + "tab": "\t", namespace + "br": "\n",
                  namespace + "cr": "\n", namespace + "p": "\r"}
    parts = []
    with ZipFile(path) as archive, archive.open("word/document.xml") as document:
        for _, element in ElementTree.iterparse(document, events=("end",)):
            if element.tag == namespace + "t":
                parts.append(element.text or "")
            elif element.tag in separators:
                parts.append(separators[element.tag])
            element.clear()
    return "".join(parts)


def select_sheet(workbook: Any, sheet: str | None = None) -> Any:
    if not sheet:
        return workbook.worksheets[0]
    for candidate in (sheet, sheet.strip()):
        if candidate in workbook.sheetnames:
            return workbook[candidate]
    raise ValueError(f"工作表「{sheet}」不存在；当前文件包含的工作表：{workbook.sheetnames}")


def list_excel_sheets(path: Path) -> dict:
    from openpyxl import load_workbook
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            return {"path": str(path), "sheets": [
                {"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column}
                for ws in workbook.worksheets]}
        finally:
            workbook.close()


def read_ppt_file(path: Path) -> str:
    # Follow presentation relationships, not filenames: slide IDs can be reordered.
    import posixpath
    relns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    docrel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    draw = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    pres = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

    def relationships(archive, part):
        folder, name = posixpath.split(part)
        relfile = posixpath.join(folder, "_rels", name + ".rels")
        if relfile not in archive.namelist():
            return {}
        result = {}
        for rel in ElementTree.fromstring(archive.read(relfile)).findall(relns + "Relationship"):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            target = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                       else posixpath.join(folder, target))
            if not target.startswith("ppt/") or "\\" in target:
                raise ValueError("PPT 包含无效的内部关系路径")
            result[rel.get("Id")] = (target, rel.get("Type", ""))
        return result

    def paragraphs(root, *, notes=False):
        result = []
        for shape in root.iter(pres + "sp") if notes else [root]:
            if notes and any(p.get("type") in {"sldNum", "hdr", "ftr", "dt", "sldImg"}
                             for p in shape.iter(pres + "ph")):
                continue
            for paragraph in shape.iter(draw + "p"):
                text = "".join(node.text or "" for node in paragraph.iter(draw + "t"))
                if text:
                    result.append(text)
        return result

    with ZipFile(path) as archive:
        main = "ppt/presentation.xml"
        links = relationships(archive, main)
        pages = []
        for index, slide in enumerate(ElementTree.fromstring(archive.read(main)).iter(pres + "sldId"), 1):
            part = links[slide.get(docrel + "id")][0]
            text = paragraphs(ElementTree.fromstring(archive.read(part)))
            for target, kind in relationships(archive, part).values():
                if kind.endswith("/notesSlide"):
                    notes = paragraphs(ElementTree.fromstring(archive.read(target)), notes=True)
                    if notes:
                        text.extend(["【备注】", *notes])
            pages.append(f"【第 {index} 页】\n" + "\n".join(text))
        return "\n\n".join(pages)
