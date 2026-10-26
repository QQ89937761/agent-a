"""Word 默认排版规范(来自用户长期记忆「Word 默认排版规范」)。

做 Word 文件时**默认套用**这套版式,除非用户明确说"这份不用记忆里的排版"。

规范要点:
- A4;左右上下页边距 25mm;页眉顶端距 15mm,页脚底端距 17.5mm;左侧装订。
- 正文:小四号宋体(英文/数字 Times New Roman),首行缩进 2 字符,
  固定值行距 22 磅,段前段后 0。
- 章标题:小三号宋体居中加粗;节标题:四号宋体加粗无缩进;条款:小四号宋体加粗无缩进。
- 表格:标题在表上方("表1…",五号黑体居中);表内五号宋体单倍行距;
  外框 1.5 磅黑线、内框细黑线;表头浅蓝底黑字,表体不加底色;表格铺满版心。
- 图标题在图下方("图1…",五号黑体居中)。
- 页眉:项目名称左对齐(五号宋体),资料不足时不加占位;页脚:页码居中。
"""
from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips

PAGE_MARGIN_MM = 25.0
HEADER_DISTANCE_MM = 15.0
FOOTER_DISTANCE_MM = 17.5

CN_FONT = "宋体"
EN_FONT = "Times New Roman"
HEI_FONT = "黑体"

SIZE_BODY = 12.0     # 小四
SIZE_CHAPTER = 15.0  # 小三
SIZE_SECTION = 14.0  # 四号
SIZE_TABLE = 10.5    # 五号
SIZE_HEADER = 10.5   # 五号

BODY_LINE_SPACING_PT = 22.0
HEADER_FILL = "DCE6F1"  # 表头浅蓝底


def _set_run_font(run, *, cn: str = CN_FONT, en: str = EN_FONT, size: float, bold: bool = False) -> None:
    """python-docx 设中文字体必须同时写 w:eastAsia,否则 Word/WPS 不认。"""
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)


def apply_default_style(doc) -> None:
    """页面尺寸、页边距、页眉页脚距离、Normal 样式。"""
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(PAGE_MARGIN_MM)
    section.bottom_margin = Mm(PAGE_MARGIN_MM)
    section.left_margin = Mm(PAGE_MARGIN_MM)
    section.right_margin = Mm(PAGE_MARGIN_MM)
    section.header_distance = Mm(HEADER_DISTANCE_MM)
    section.footer_distance = Mm(FOOTER_DISTANCE_MM)
    section.gutter = Mm(0)

    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(SIZE_BODY)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)


def _body_format(paragraph, *, indent: bool, bold: bool = False, align=None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(BODY_LINE_SPACING_PT)
    if indent:
        # 首行缩进 2 字符:按正文字号折算(小四 12pt → 24pt)
        pf.first_line_indent = Pt(SIZE_BODY * 2)
    else:
        pf.first_line_indent = Pt(0)
    if align is not None:
        paragraph.alignment = align


def add_chapter(doc, text: str):
    """章标题:小三号宋体,居中,加粗。"""
    para = doc.add_paragraph()
    _body_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run_font(para.add_run(text), size=SIZE_CHAPTER, bold=True)
    return para


def add_section(doc, text: str):
    """节标题:四号宋体,加粗,无缩进。"""
    para = doc.add_paragraph()
    _body_format(para, indent=False)
    _set_run_font(para.add_run(text), size=SIZE_SECTION, bold=True)
    return para


def add_body(doc, text: str):
    """正文:小四号宋体,首行缩进 2 字符,固定值行距 22 磅。"""
    para = doc.add_paragraph()
    _body_format(para, indent=True)
    _set_run_font(para.add_run(text), size=SIZE_BODY)
    return para


def add_clause(doc, text: str):
    """条款:小四号宋体,加粗,无缩进。"""
    para = doc.add_paragraph()
    _body_format(para, indent=False)
    _set_run_font(para.add_run(text), size=SIZE_BODY, bold=True)
    return para


def _cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _table_borders(table) -> None:
    """外框 1.5 磅黑线,内框细黑线。"""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", 12), ("left", 12), ("bottom", 12), ("right", 12),
                       ("insideH", 4), ("insideV", 4)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))  # 1/8 磅:12 → 1.5 磅
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tbl_pr.append(borders)


def _set_table_width(table, width_twips: int) -> None:
    """固定表格总宽度并平均分配列宽,避免只按内容自动收缩。"""
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(width_twips))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    column_count = len(table.columns)
    if not column_count:
        return
    base_width, remainder = divmod(width_twips, column_count)
    for index, column in enumerate(table.columns):
        column_width = base_width + (1 if index < remainder else 0)
        width = Twips(column_width)
        column.width = width
        for cell in column.cells:
            cell.width = width
        grid_column = table._tbl.tblGrid.gridCol_lst[index]
        grid_column.set(qn("w:w"), str(column_width))


def add_table(doc, headers, rows, *, caption: str = ""):
    """表格:标题在表上方(五号黑体居中),表内五号宋体,表头浅蓝底黑字。"""
    headers = list(headers or [])
    rows = [list(row) for row in (rows or [])]
    if not headers and not rows:
        return None
    width = max([len(headers)] + [len(row) for row in rows]) if (headers or rows) else 0
    if width <= 0:
        return None
    if caption:
        para = doc.add_paragraph()
        _body_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_run_font(para.add_run(caption), cn=HEI_FONT, en=HEI_FONT, size=SIZE_TABLE)
    table = doc.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    section = doc.sections[-1]
    text_width_emu = section.page_width - section.left_margin - section.right_margin - section.gutter
    _set_table_width(table, max(1, int(round(text_width_emu / 635))))
    _table_borders(table)
    if headers:
        cells = table.add_row().cells
        for index in range(width):
            value = headers[index] if index < len(headers) else ""
            cells[index].text = ""
            para = cells[index].paragraphs[0]
            _body_format(para, indent=False)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            _set_run_font(para.add_run(str(value)), size=SIZE_TABLE, bold=True)
            _cell_shading(cells[index], HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for index in range(width):
            value = row[index] if index < len(row) else ""
            cells[index].text = ""
            para = cells[index].paragraphs[0]
            _body_format(para, indent=False)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            _set_run_font(para.add_run(str(value)), size=SIZE_TABLE)
    return table


def add_figure_caption(doc, text: str):
    """图标题:图形下方,五号黑体,居中。"""
    para = doc.add_paragraph()
    _body_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run_font(para.add_run(text), cn=HEI_FONT, en=HEI_FONT, size=SIZE_TABLE)
    return para


def set_header(doc, text: str) -> None:
    """页眉:项目名称左对齐,五号宋体;空文本时留空(不加占位)。"""
    header = doc.sections[0].header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if text:
        _set_run_font(para.add_run(text), size=SIZE_HEADER)


def add_page_number_footer(doc) -> None:
    """页脚正中插入页码。"""
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    footer = doc.sections[0].footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    _set_run_font(run, size=SIZE_HEADER)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._element.append(begin)
    run._element.append(instr)
    run._element.append(separate)
    run._element.append(result)
    run._element.append(end)
