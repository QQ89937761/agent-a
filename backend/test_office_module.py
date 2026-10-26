"""Office COM ownership regression tests without requiring Office itself."""
from __future__ import annotations

import unittest
import asyncio
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from types import ModuleType
from unittest import mock
import sys

from backend.core.tools import ToolRegistry
from backend.modules.office import office_module


class _Documents:
    def __init__(self, doc):
        self._doc = doc
        self.Count = 1

    def Item(self, index):
        assert index == 1
        return self._doc

    def Open(self, _path, **kwargs):
        self.open_kwargs = kwargs
        return self._doc


class _Find:
    def __init__(self, doc):
        self._doc = doc

    def Execute(self, find, *_args):
        replace = _args[-2]
        self._doc.Content.Text = self._doc.Content.Text.replace(find, replace)
        self._doc.Saved = False
        return True


class _Content:
    def __init__(self, doc, text: str):
        self.Text = text
        self.Find = _Find(doc)


class _Document:
    def __init__(self, path: Path, *, text: str = "", fail_close: bool = False):
        self.FullName = str(path)
        self.fail_close = fail_close
        self.close_called = False
        self.save_called = False
        self.disk_text = text
        self.Saved = False
        self.ReadOnly = False
        self.Content = _Content(self, text)

    def Close(self, **_kwargs):
        self.close_called = True
        if self.fail_close:
            raise RuntimeError("WPS already exited")

    def Save(self):
        self.save_called = True
        self.disk_text = self.Content.Text
        self.Saved = True


class _Application:
    def __init__(self, doc, *, fail_quit: bool = False):
        self.Documents = _Documents(doc)
        self.fail_quit = fail_quit
        self.quit_called = False
        self.Visible = True
        self.DisplayAlerts = True

    def Quit(self):
        self.quit_called = True
        if self.fail_quit:
            raise RuntimeError("RPC server unavailable")


def _cell_address(row: int, column: int) -> str:
    return f"{chr(64 + column)}{row}"


class _ExcelRange:
    def __init__(self, state, address: str):
        if address == "INVALID":
            raise ValueError("invalid range")
        self._state = state
        self._address = address

    @property
    def Value(self):
        if ":" not in self._address:
            return self._state.get(self._address)
        start, end = self._address.split(":", 1)
        start_col, start_row = ord(start[0]) - 64, int(start[1:])
        end_col, end_row = ord(end[0]) - 64, int(end[1:])
        return tuple(
            tuple(
                self._state.get(_cell_address(row, column))
                for column in range(start_col, end_col + 1)
            )
            for row in range(start_row, end_row + 1)
        )

    @Value.setter
    def Value(self, value):
        self._state[self._address] = value


class _ExcelUsedRange:
    def __init__(self, state):
        rows = [int(address[1:]) for address in state] or [1]
        columns = [ord(address[0]) - 64 for address in state] or [1]
        self.Row = min(rows)
        self.Rows = mock.Mock(Count=max(rows) - self.Row + 1)
        self.Value = tuple(
            tuple(state.get(_cell_address(row, column)) for column in range(1, max(columns) + 1))
            for row in range(self.Row, max(rows) + 1)
        )


class _ExcelWorksheet:
    def __init__(self, state):
        self._state = state

    def Range(self, address):
        return _ExcelRange(self._state, address)

    @property
    def UsedRange(self):
        return _ExcelUsedRange(self._state)

    def Cells(self, row, column):
        return _ExcelRange(self._state, _cell_address(row, column))


class _ExcelWorksheets:
    def __init__(self, state):
        self._sheet = _ExcelWorksheet(state)

    def __call__(self, _key):
        return self._sheet


class _ExcelWorkbook:
    def __init__(self, path: Path, state):
        self.FullName = str(path)
        self.Worksheets = _ExcelWorksheets(state)
        self.Saved = False
        self.ReadOnly = False
        self.save_called = False
        self.close_called = False

    def Save(self):
        self.Saved = True
        self.save_called = True

    def SaveAs(self, path):
        self.FullName = str(path)
        self.Save()

    def Close(self, **_kwargs):
        self.close_called = True


class _ExcelWorkbooks:
    def __init__(self, states, paths=()):
        self._states = states
        self._items = [
            _ExcelWorkbook(path, states.setdefault(Path(path).resolve(), {})) for path in paths
        ]

    @property
    def Count(self):
        return len(self._items)

    def Item(self, index):
        return self._items[index - 1]

    def Open(self, path, **kwargs):
        self.open_kwargs = kwargs
        self.open_kwargs = kwargs
        resolved = Path(path).resolve()
        workbook = _ExcelWorkbook(resolved, self._states.setdefault(resolved, {}))
        self._items.append(workbook)
        return workbook

    def Add(self):
        workbook = _ExcelWorkbook(Path("new.xlsx"), {})
        self._items.append(workbook)
        return workbook


class _ExcelApplication:
    def __init__(self, states, paths=()):
        self.Workbooks = _ExcelWorkbooks(states, paths)
        self.Visible = True
        self.DisplayAlerts = True
        self.quit_called = False

    def Quit(self):
        self.quit_called = True


class _ExcelClient:
    def __init__(self, states, active=None):
        self.states = states
        self.active = active
        self.owned_apps = []

    def GetActiveObject(self, progid):
        if self.active is None:
            raise RuntimeError(f"no active {progid}")
        return self.active

    def DispatchEx(self, progid):
        if progid != "Excel.Application":
            raise AssertionError(progid)
        app = _ExcelApplication(self.states)
        self.owned_apps.append(app)
        return app


@contextmanager
def _registered_tools(client=None):
    if client is None:
        with mock.patch.object(office_module, "_com_available", return_value=False):
            registry = ToolRegistry()
            office_module.register_tools(registry)
            yield registry
        return
    client_module = ModuleType("win32com.client")
    client_module.GetActiveObject = client.GetActiveObject
    client_module.DispatchEx = client.DispatchEx
    win32com_module = ModuleType("win32com")
    win32com_module.__path__ = []
    win32com_module.client = client_module
    pythoncom_module = ModuleType("pythoncom")
    pythoncom_module.CoInitialize = lambda: None
    pythoncom_module.CoUninitialize = lambda: None
    modules = {
        "win32com": win32com_module,
        "win32com.client": client_module,
        "pythoncom": pythoncom_module,
    }
    with mock.patch.dict(sys.modules, modules), mock.patch.object(
        office_module, "_com_available", return_value=True
    ):
        registry = ToolRegistry()
        office_module.register_tools(registry)
        yield registry


class OfficeFileReadingTests(unittest.TestCase):
    def test_excel_empty_single_cell_and_invalid_ranges(self):
        from openpyxl import Workbook
        from backend.modules.office.office_reading import read_excel_file
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.xlsx"
            workbook = Workbook()
            workbook.save(path)
            self.assertIsNone(read_excel_file(path))
            self.assertIsNone(read_excel_file(path, cell_range="A1"))
            workbook.active["B2"] = 42
            workbook.save(path)
            self.assertEqual(42, read_excel_file(path))
            self.assertEqual(42, read_excel_file(path, cell_range="B2"))
            self.assertEqual(((None, None), (None, 42)), read_excel_file(path, cell_range="A1:B2"))
            for area in ("A:A", "1:2", "A0", "B2:A1", "XFE1", "A1048577", "INVALID"):
                with self.subTest(area=area), self.assertRaises(ValueError):
                    read_excel_file(path, cell_range=area)

    def test_xlsm_returns_formula_cache_without_running_formula_or_macro(self):
        from io import BytesIO
        from zipfile import ZipFile
        from openpyxl import Workbook
        from backend.modules.office.office_reading import read_excel_file
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "formula.xlsm"
            workbook = Workbook()
            workbook.active["A1"] = "=1+1"
            workbook.active["B1"] = '=WEBSERVICE("https://example.test")'
            workbook.save(path)
            data = path.read_bytes()
            with ZipFile(BytesIO(data)) as source, ZipFile(path, "w") as target:
                for member in source.infolist():
                    content = source.read(member.filename)
                    if member.filename == "xl/worksheets/sheet1.xml":
                        content = content.replace(b"<v></v>", b"<v>2</v>", 1)
                    target.writestr(member, content)
                target.writestr("xl/vbaProject.bin", b"not executable: macro content is never loaded")
            self.assertEqual(((2, None),), read_excel_file(path, cell_range="A1:B1"))


class OfficeWordLifecycleTests(unittest.TestCase):
    def test_unreachable_owner_file_is_not_assumed_stale(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "closed.docx").resolve()
            path.touch()
            path.with_name("~$osed.docx").touch()
            doc = _Document(path)
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active WPS")
            client.DispatchEx.return_value = app

            clock = 0.0

            def monotonic():
                nonlocal clock
                clock += 1.0
                return clock

            with mock.patch.object(office_module.time, "monotonic", side_effect=monotonic), \
                    mock.patch.object(office_module.time, "sleep"):
                with self.assertRaisesRegex(RuntimeError, "占用标记"):
                    office_module._open_word_document(client, path, read_only=True)

            client.DispatchEx.assert_not_called()

    def test_short_filename_does_not_treat_old_malformed_owner_name_as_real(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "ab.docx").resolve()
            path.touch()
            path.with_name("~$.docx").touch()
            doc = _Document(path)
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active WPS")
            client.DispatchEx.return_value = app

            with mock.patch.object(
                office_module.time, "monotonic", side_effect=AssertionError("must not wait")
            ):
                office_module._open_word_document(client, path, read_only=True)

            client.DispatchEx.assert_called_once_with("Word.Application")

    def test_open_document_is_reused_without_closing_user_wps(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "open.docx").resolve()
            path.touch()
            doc = _Document(path)
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.return_value = app

            word, opened, owned = office_module._open_word_document(client, path)
            office_module._release_word_document(word, opened, owned)

            self.assertFalse(owned)
            self.assertFalse(doc.close_called)
            self.assertFalse(app.quit_called)
            client.DispatchEx.assert_not_called()

    def test_owned_wps_cleanup_errors_do_not_mask_completed_operation(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "closed.docx").resolve()
            path.touch()
            doc = _Document(path, fail_close=True)
            app = _Application(doc, fail_quit=True)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active WPS")
            client.DispatchEx.return_value = app

            word, opened, owned = office_module._open_word_document(client, path)
            office_module._release_word_document(word, opened, owned)

            self.assertTrue(owned)
            self.assertTrue(doc.close_called)
            self.assertTrue(app.quit_called)


class OfficeToolOwnershipTests(unittest.IsolatedAsyncioTestCase):
    async def test_inaccessible_wps_owner_blocks_writes_and_legacy_reads(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            cases = (
                ("excel_write_cells", "book.xlsx", {"cells": {"A1": "changed"}}),
                ("excel_append_row", "book.xlsx", {"values": ["changed"]}),
                ("word_find_replace", "draft.docx", {"find": "old", "replace": "new"}),
                ("excel_read_range", "legacy.xls", {"cell_range": "A1"}),
                ("word_read_text", "legacy.doc", {}),
            )
            for name, filename, extra in cases:
                path = root / filename
                path.write_bytes(b"saved file must remain untouched")
                path.with_name(f"~${path.name}").touch()
                client = mock.Mock()
                client.GetActiveObject.side_effect = RuntimeError("WPS not in ROT")
                with self.subTest(name=name), _registered_tools(client) as registry:
                    with mock.patch.object(office_module, "time") as clock:
                        clock.monotonic.side_effect = [0.0, 3.0]
                        with self.assertRaisesRegex(RuntimeError, "占用标记"):
                            await registry.call(name, {"path": str(path), **extra})
                client.DispatchEx.assert_not_called()
                self.assertEqual(b"saved file must remain untouched", path.read_bytes())

    async def test_saved_reads_with_owner_marker_warn_with_and_without_com(self):
        from openpyxl import Workbook
        from docx import Document
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            excel = root / "book.xlsx"
            wb = Workbook()
            wb.active["A1"] = "已保存"
            wb.save(excel)
            word = root / "draft.docx"
            doc = Document()
            doc.add_paragraph("已保存正文")
            doc.save(word)
            for path in (excel, word):
                path.with_name(f"~${path.name}").touch()
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no ROT")
            for com in (None, client):
                with _registered_tools(com) as registry:
                    for name, extra in (("excel_read_all", {}), ("excel_read_range", {"cell_range": "A1"})):
                        result = await registry.call(name, {"path": str(excel), **extra})
                        self.assertIn("磁盘已保存版本", result["warning"])
                        self.assertEqual("已保存", result["data"])
                    result = await registry.call("word_read_text", {"path": str(word)})
                    self.assertIn("磁盘已保存版本", result)
                    self.assertIn("已保存正文", result)
            client.DispatchEx.assert_not_called()

    async def test_readonly_com_write_and_silent_save_failure_never_report_success(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "book.xls").resolve()
            path.touch()
            for name, extra in (
                ("excel_write_cells", {"cells": {"A1": "changed"}}),
                ("excel_append_row", {"values": ["changed"]}),
            ):
                for readonly in (True, False):
                    states = {path: {"A1": "original"}}
                    app = _ExcelApplication(states, [path])
                    workbook = app.Workbooks.Item(1)
                    app.Workbooks.Open = mock.Mock(return_value=workbook)
                    workbook.ReadOnly = readonly
                    workbook.Save = mock.Mock()  # Silent failure leaves Saved=False.
                    client = mock.Mock()
                    client.GetActiveObject.side_effect = RuntimeError("no active Excel")
                    client.DispatchEx.return_value = app
                    with self.subTest(name=name, readonly=readonly), _registered_tools(client) as registry:
                        with self.assertRaisesRegex(RuntimeError, "只读" if readonly else "未确认保存成功"):
                            await registry.call(name, {"path": str(path), **extra})
                    if readonly:
                        self.assertEqual({"A1": "original"}, states[path])
                        workbook.Save.assert_not_called()
                    self.assertTrue(workbook.close_called)
                    self.assertTrue(app.quit_called)

            path = (Path(tmp) / "draft.doc").resolve()
            path.touch()
            for readonly in (True, False):
                doc = _Document(path, text="old")
                doc.ReadOnly = readonly
                doc.Save = mock.Mock()
                app = _Application(doc)
                client = mock.Mock()
                client.GetActiveObject.side_effect = RuntimeError("no active Word")
                client.DispatchEx.return_value = app
                with self.subTest(word_readonly=readonly), _registered_tools(client) as registry:
                    with self.assertRaisesRegex(RuntimeError, "只读" if readonly else "未确认保存成功"):
                        await registry.call("word_find_replace", {
                            "path": str(path), "find": "old", "replace": "new"
                        })
                if readonly:
                    self.assertEqual("old", doc.Content.Text)
                    doc.Save.assert_not_called()
                self.assertTrue(doc.close_called)
                self.assertTrue(app.quit_called)

    async def test_registration_guard_exposes_every_office_tool(self):
        client = _ExcelClient({})
        with _registered_tools(client) as registry:
            self.assertEqual(set(office_module.EXPOSED_TOOLS), set(registry.names()))

    async def test_all_com_writers_reject_system_paths_before_opening_office(self):
        for path in ('C:/Windows/Temp/office.xlsx', 'C:/Program Files/office.docx', 'C:/ProgramData/office.xlsx'):
            for name, args in (
                ('excel_write_cells', {'cells': {'A1': 'value'}}),
                ('excel_append_row', {'values': ['value']}),
                ('word_find_replace', {'find': 'old', 'replace': 'new'}),
            ):
                with self.subTest(path=path, name=name):
                    client = mock.Mock()
                    with _registered_tools(client) as registry:
                        with self.assertRaisesRegex(PermissionError, '系统目录'):
                            await registry.call(name, {'path': path, **args})
                    client.GetActiveObject.assert_not_called()
                    client.DispatchEx.assert_not_called()

    async def test_office_writer_descriptions_share_once_per_conversation_notice(self):
        with _registered_tools(_ExcelClient({})) as registry:
            for name in ('excel_write_cells', 'excel_append_row', 'word_find_replace'):
                tool = registry.get(name)
                self.assertTrue(tool.side_effect)
                self.assertTrue(tool.danger)
                self.assertEqual(1, tool.description.count('同一对话最多请求一次危险操作授权'))
                self.assertIn('已有宏、外链或事件', tool.description)

    async def test_common_saved_office_files_read_without_starting_com(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            xlsx = root / "saved.xlsx"
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "数据"
            sheet["A1"] = "已保存"
            sheet["B1"] = 7
            sheet["C1"] = "=1+1"
            workbook.calculation.fullCalcOnLoad = False
            workbook.save(xlsx)
            docx = root / "saved.docx"
            from docx import Document
            document = Document()
            document.add_paragraph("第一段")
            document.add_paragraph("第二段")
            document.save(docx)

            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active Office")
            for com_client in (None, client):
                with self.subTest(com_available=com_client is not None), _registered_tools(com_client) as registry:
                    expected = office_module.EXPOSED_TOOLS
                    self.assertEqual(set(expected), set(registry.names()))
                    values = await registry.call("excel_read_range", {
                        "path": str(xlsx), "sheet": "数据", "cell_range": "A1:C1"
                    })
                    all_values = await registry.call("excel_read_all", {"path": str(xlsx), "sheet": "数据"})
                    text = await registry.call("word_read_text", {"path": str(docx)})
                    self.assertEqual((("已保存", 7, None),), values)
                    self.assertEqual((("已保存", 7),), all_values)
                    self.assertIn("第一段", text)
                    self.assertIn("第二段", text)
            client.DispatchEx.assert_not_called()

    async def test_common_format_aliases_never_fall_back_to_com(self):
        # 模拟 symlink/路径映射，无需 Windows 创建链接权限。
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            cases = (
                ("excel_read_range", "alias.xlsx", "real.xls", {"cell_range": "A1"}),
                ("excel_read_all", "alias.xlsm", "real.xls", {}),
                ("word_read_text", "alias.docm", "real.doc", {}),
            )
            for name, alias, target, extra in cases:
                with self.subTest(name=name):
                    resolved = root / target
                    resolved.write_bytes(b"legacy or corrupt data")
                    client = mock.Mock()
                    client.GetActiveObject.side_effect = RuntimeError("no active Office")
                    with _registered_tools(client) as registry, mock.patch.object(
                        office_module, "_resolve", return_value=resolved
                    ):
                        from zipfile import BadZipFile
                        with self.assertRaises(BadZipFile):
                            await registry.call(name, {"path": str(root / alias), **extra})
                    client.DispatchEx.assert_not_called()

    async def test_excel_alias_with_different_target_suffix_still_reads_as_data(self):
        with TemporaryDirectory() as tmp:
            from openpyxl import Workbook
            root = Path(tmp)
            target = root / "real.xls"
            workbook = Workbook()
            workbook.active["A1"] = "safe OOXML data"
            workbook.save(target)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active Excel")
            with _registered_tools(client) as registry, mock.patch.object(
                office_module, "_resolve", return_value=target
            ):
                for name, extra in (("excel_read_range", {"cell_range": "A1"}), ("excel_read_all", {})):
                    self.assertEqual("safe OOXML data", await registry.call(
                        name, {"path": str(root / "alias.xlsx"), **extra}
                    ))
            client.DispatchEx.assert_not_called()

    async def test_corrupt_common_files_and_missing_sheets_do_not_start_com(self):
        from zipfile import BadZipFile
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active Office")
            with _registered_tools(client) as registry:
                for name, filename, extra in (
                    ("excel_read_range", "broken.xlsx", {"cell_range": "A1"}),
                    ("excel_read_all", "broken.xlsm", {}),
                    ("word_read_text", "broken.docx", {}),
                ):
                    path = root / filename
                    path.write_bytes(b"not a zip")
                    with self.subTest(name=name), self.assertRaises(BadZipFile):
                        await registry.call(name, {"path": str(path), **extra})
                from openpyxl import Workbook
                path = root / "valid.xlsx"
                Workbook().save(path)
                with self.assertRaisesRegex(ValueError, "工作表.*Sheet"):
                    await registry.call("excel_read_all", {"path": str(path), "sheet": "missing"})
            client.DispatchEx.assert_not_called()

    async def test_legacy_reads_without_com_report_requirement(self):
        with _registered_tools() as registry:
            self.assertEqual(set(office_module.EXPOSED_TOOLS), set(registry.names()))
            for name, path, extra in (
                ("excel_read_range", "saved.xls", {"cell_range": "A1"}),
                ("excel_read_all", "saved.xls", {}),
                ("word_read_text", "saved.doc", {}),
            ):
                with self.subTest(name=name), self.assertRaisesRegex(RuntimeError, "COM"):
                    await registry.call(name, {"path": path, **extra})

    async def test_saved_docm_preserves_tables_tabs_breaks_and_pagination(self):
        from docx import Document
        from backend.modules.office.office_reading import read_word_file
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "saved.docm"
            document = Document()
            document.add_paragraph("正文")
            run = document.add_paragraph().add_run("左")
            run.add_tab()
            run.add_text("右")
            run.add_break()
            run.add_text("下一行")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "表格一"
            table.cell(0, 1).text = "表格二"
            document.add_paragraph("中" * 21000)
            document.save(path)
            expected = read_word_file(path)
            self.assertTrue(expected.startswith("正文\r左\t右\n下一行\r表格一\r表格二\r"))
            with _registered_tools() as registry:
                first = await registry.call("word_read_text", {"path": str(path)})
                second = await registry.call("word_read_text", {"path": str(path), "offset": 20000})
                self.assertIn("下一页 offset=20000", first)
                self.assertIn("状态=已读完", second)
                self.assertEqual(expected, first.split("\n", 1)[1] + second.split("\n", 1)[1])
                for offset in (-1, True, 1.5, len(expected) + 1):
                    with self.subTest(offset=offset), self.assertRaises(ValueError):
                        await registry.call("word_read_text", {"path": str(path), "offset": offset})

    async def test_read_open_word_returns_unsaved_text_without_save_or_close(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "open.docx").resolve()
            path.write_bytes(b"disk deliberately not parseable")
            doc = _Document(path, text="saved")
            doc.Content.Text = "current unsaved draft"
            path.with_name(f"~${path.name}").touch()
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.return_value = app
            with _registered_tools(client) as registry:
                result = await registry.call("word_read_text", {"path": str(path)})
            self.assertEqual("current unsaved draft", result.split("\n", 1)[1])
            self.assertEqual("saved", doc.disk_text)
            self.assertFalse(doc.save_called)
            self.assertFalse(doc.close_called)
            self.assertFalse(app.quit_called)
            client.DispatchEx.assert_not_called()

    async def test_legacy_excel_reads_use_read_only_and_disable_link_updates(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "legacy.xls").resolve()
            path.touch()
            client = _ExcelClient({path: {"A1": "legacy value"}})
            with _registered_tools(client) as registry:
                self.assertEqual("legacy value", await registry.call(
                    "excel_read_range", {"path": str(path), "cell_range": "A1"}
                ))
                self.assertEqual((("legacy value",),), await registry.call(
                    "excel_read_all", {"path": str(path)}
                ))
            for app in client.owned_apps:
                self.assertEqual({"ReadOnly": True, "UpdateLinks": 0}, app.Workbooks.open_kwargs)
                self.assertTrue(app.Workbooks.Item(1).close_called)
                self.assertTrue(app.quit_called)

    async def test_saved_office_reads_are_not_dangerous_but_legacy_formats_are(self):
        from backend.core.agent import _looks_dangerous, _remote_confirmation_categories
        for name, path in (
            ("excel_read_range", "D:/work/report.xlsx"),
            ("excel_read_all", "D:/work/report.xlsm"),
            ("word_read_text", "D:/work/report.docm"),
        ):
            with self.subTest(name=name):
                args = {"path": path}
                self.assertFalse(_looks_dangerous(name, args))
                self.assertEqual(set(), _remote_confirmation_categories(name, args))
        for name, path in (
            ("excel_read_range", "D:/work/report.xls"),
            ("excel_read_all", "D:/work/report.ods"),
            ("word_read_text", "D:/work/report.doc"),
        ):
            with self.subTest(name=name):
                args = {"path": path}
                self.assertTrue(_looks_dangerous(name, args))
                self.assertEqual({"external_code"}, _remote_confirmation_categories(name, args))

    async def test_word_read_text_pages_without_saving_and_cleans_up(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "long.doc").resolve()
            path.touch()
            original = "a" * 20000 + "中" * 20000 + "tail"
            doc = _Document(path, text=original)
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active Word")
            client.DispatchEx.return_value = app

            with _registered_tools(client) as registry:
                first = await registry.call("word_read_text", {"path": str(path)})
                self.assertIn("状态=未读完；请继续读取，下一页 offset=20000", first)
                second = await registry.call("word_read_text", {"path": str(path), "offset": 20000})
                self.assertIn("状态=未读完；请继续读取，下一页 offset=40000", second)
                third = await registry.call("word_read_text", {"path": str(path), "offset": 40000})
                self.assertIn("状态=已读完", third)

            self.assertEqual(original, "".join(
                page.split("\n", 1)[1] for page in (first, second, third)
            ))
            self.assertEqual(original, doc.Content.Text)
            self.assertEqual(original, doc.disk_text)
            self.assertFalse(doc.save_called)
            self.assertTrue(doc.close_called)
            self.assertTrue(app.quit_called)
            self.assertEqual({"ReadOnly": True}, app.Documents.open_kwargs)

            failed_client = mock.Mock()
            failed_client.GetActiveObject.side_effect = RuntimeError("no active Word")
            failed_client.DispatchEx.side_effect = RuntimeError("Word unavailable")
            with _registered_tools(failed_client) as registry:
                with self.assertRaisesRegex(RuntimeError, "Word unavailable"):
                    await registry.call("word_read_text", {"path": str(path)})

    async def test_word_find_replace_saves_only_owned_document(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "owned.doc").resolve()
            path.touch()
            doc = _Document(path, text="old value and old value")
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.side_effect = RuntimeError("no active Word")
            client.DispatchEx.return_value = app

            with _registered_tools(client) as registry:
                result = await registry.call(
                    "word_find_replace",
                    {"path": str(path), "find": "old", "replace": "new"},
                )

            self.assertIn("old", result)
            self.assertEqual("new value and new value", doc.Content.Text)
            self.assertEqual(doc.Content.Text, doc.disk_text)
            self.assertTrue(doc.save_called)
            self.assertTrue(doc.close_called)
            self.assertTrue(app.quit_called)

    async def test_excel_read_all_preserves_user_open_workbook(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "open.xlsx").resolve()
            path.touch()
            states = {path: {"A1": "name", "B1": "count", "A2": "first", "B2": 1}}
            active = _ExcelApplication(states, [path])
            workbook = active.Workbooks.Item(1)
            client = _ExcelClient(states, active=active)

            with _registered_tools(client) as registry:
                values = await registry.call("excel_read_all", {"path": str(path)})

            self.assertEqual((("name", "count"), ("first", 1)), values)
            self.assertEqual(1, states[path]["B2"])
            self.assertFalse(workbook.save_called)
            self.assertFalse(workbook.close_called)
            self.assertFalse(active.quit_called)
            self.assertEqual([], client.owned_apps)

            failed_client = mock.Mock()
            failed_client.GetActiveObject.side_effect = RuntimeError("no active Excel")
            failed_client.DispatchEx.side_effect = RuntimeError("Excel unavailable")
            with _registered_tools(failed_client) as registry:
                with self.assertRaisesRegex(RuntimeError, "Excel unavailable"):
                    await registry.call("excel_read_all", {"path": str(path.with_suffix(".xls"))})

    async def test_excel_append_row_saves_owned_file_and_rejects_open_target(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "append.xls").resolve()
            path.touch()
            states = {path: {"A1": "name", "B1": "count", "A2": "first", "B2": 1}}
            client = _ExcelClient(states)

            with _registered_tools(client) as registry:
                result = await registry.call(
                    "excel_append_row", {"path": str(path), "values": ["last", 2]}
                )

            workbook = client.owned_apps[0].Workbooks.Item(1)
            self.assertIn("第 3 行", result)
            self.assertEqual("last", states[path]["A3"])
            self.assertEqual(2, states[path]["B3"])
            self.assertTrue(workbook.save_called)
            self.assertTrue(workbook.close_called)
            self.assertTrue(client.owned_apps[0].quit_called)

            open_states = {path: {"A1": "unchanged"}}
            active = _ExcelApplication(open_states, [path])
            open_workbook = active.Workbooks.Item(1)
            open_client = _ExcelClient(open_states, active=active)
            with _registered_tools(open_client) as registry:
                with self.assertRaisesRegex(RuntimeError, "请先确认并保存后关闭"):
                    await registry.call(
                        "excel_append_row", {"path": str(path), "values": ["blocked"]}
                    )

            self.assertEqual({"A1": "unchanged"}, open_states[path])
            self.assertFalse(open_workbook.save_called)
            self.assertFalse(open_workbook.close_called)
            self.assertFalse(active.quit_called)

    async def test_all_four_excel_tools_use_owned_hidden_noninteractive_instances(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "owned.xls").resolve()
            path.touch()
            states = {path: {"A1": "name", "B1": "count", "A2": "first", "B2": 1}}
            client = _ExcelClient(states)

            with _registered_tools(client) as registry:
                await registry.call(
                    "excel_write_cells", {"path": str(path), "cells": {"B2": 2}}
                )
                values = await registry.call(
                    "excel_read_range", {"path": str(path), "cell_range": "A1:B2"}
                )
                all_values = await registry.call("excel_read_all", {"path": str(path)})
                await registry.call(
                    "excel_append_row", {"path": str(path), "values": ["last", 3]}
                )

            self.assertEqual((("name", "count"), ("first", 2)), values)
            self.assertEqual(values, all_values)
            self.assertEqual("last", states[path]["A3"])
            self.assertEqual(3, states[path]["B3"])
            self.assertEqual(4, len(client.owned_apps))
            for app in client.owned_apps:
                self.assertFalse(app.Visible)
                self.assertFalse(app.DisplayAlerts)
                self.assertTrue(app.quit_called)
                self.assertTrue(app.Workbooks.Item(1).close_called)

    async def test_user_excel_with_unsaved_work_is_untouched_for_other_file(self):
        with TemporaryDirectory() as tmp:
            user_path = (Path(tmp) / "user.xls").resolve()
            target_path = (Path(tmp) / "target.xls").resolve()
            user_path.touch()
            target_path.touch()
            states = {
                user_path: {"A1": "unsaved user text"},
                target_path: {"A1": "target value"},
            }
            active = _ExcelApplication(states, [user_path])
            user_workbook = active.Workbooks.Item(1)
            client = _ExcelClient(states, active=active)

            with _registered_tools(client) as registry:
                value = await registry.call(
                    "excel_read_range", {"path": str(target_path), "cell_range": "A1"}
                )

            self.assertEqual("target value", value)
            self.assertTrue(active.Visible)
            self.assertTrue(active.DisplayAlerts)
            self.assertFalse(active.quit_called)
            self.assertFalse(user_workbook.close_called)
            self.assertFalse(user_workbook.Saved)

    async def test_open_excel_target_can_be_read_but_write_is_rejected(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "open.xlsx").resolve()
            path.touch()
            states = {path: {"A1": "unsaved"}}
            active = _ExcelApplication(states, [path])
            workbook = active.Workbooks.Item(1)
            client = _ExcelClient(states, active=active)

            with _registered_tools(client) as registry:
                value = await registry.call(
                    "excel_read_range", {"path": str(path), "cell_range": "A1"}
                )
                with self.assertRaisesRegex(RuntimeError, "请先确认并保存后关闭"):
                    await registry.call(
                        "excel_write_cells", {"path": str(path), "cells": {"A1": "changed"}}
                    )

            self.assertEqual("unsaved", value)
            self.assertEqual("unsaved", states[path]["A1"])
            self.assertFalse(workbook.save_called)
            self.assertFalse(workbook.close_called)
            self.assertFalse(active.quit_called)

    async def test_invalid_excel_range_cleans_up_without_alerts_or_hanging(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "invalid.xls").resolve()
            path.touch()
            client = _ExcelClient({path: {"A1": "value"}})

            with _registered_tools(client) as registry:
                with self.assertRaisesRegex(ValueError, "invalid range"):
                    await asyncio.wait_for(
                        registry.call(
                            "excel_read_range",
                            {"path": str(path), "cell_range": "INVALID"},
                        ),
                        timeout=1,
                    )

            app = client.owned_apps[0]
            self.assertFalse(app.DisplayAlerts)
            self.assertTrue(app.Workbooks.Item(1).close_called)
            self.assertTrue(app.quit_called)

    async def test_open_word_document_is_refused_instead_of_edited(self):
        """用户自己打开的文档一律拒绝写操作(与 Excel 同口径),不改也不存。"""
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "draft.docx").resolve()
            path.touch()
            doc = _Document(path, text="user draft and replace me")
            app = _Application(doc)
            client = mock.Mock()
            client.GetActiveObject.return_value = app

            with _registered_tools(client) as registry:
                with self.assertRaisesRegex(RuntimeError, "已打开"):
                    await registry.call(
                        "word_find_replace",
                        {"path": str(path), "find": "replace me", "replace": "changed"},
                    )

            # 拒绝发生在 Find.Execute 之前:内存与磁盘内容都必须原样
            self.assertEqual("user draft and replace me", doc.Content.Text)
            self.assertEqual("user draft and replace me", doc.disk_text)
            self.assertFalse(doc.save_called)
            self.assertFalse(doc.close_called)
            self.assertFalse(app.quit_called)

if __name__ == "__main__":
    unittest.main()
