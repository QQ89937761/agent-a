"""Behavior tests for document generation tools."""
from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock
from zipfile import ZipFile

from openpyxl import load_workbook

from backend.core.tools import ToolRegistry
from backend.modules.office import docgen_module


class DocgenToolsTests(unittest.IsolatedAsyncioTestCase):
    def test_new_documents_use_dynamic_risk_instead_of_blanket_danger(self):
        registry = ToolRegistry()
        docgen_module.register_tools(registry)
        self.assertTrue(registry.get('create_word_doc').side_effect)
        self.assertFalse(registry.get('create_word_doc').danger)
        self.assertFalse(registry.get('create_excel_file').danger)

    async def test_word_existing_target_is_preserved_without_opening(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / 'existing.docx').resolve()
            original = b'existing Word bytes'
            path.write_bytes(original)
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, '_open_file', create=True) as open_file:
                with self.assertRaisesRegex(ValueError, '避免覆盖'):
                    await registry.call('create_word_doc', {
                        'path': str(path), 'paragraphs': ['支付; format D:'],
                    })
            self.assertEqual(original, path.read_bytes())
            open_file.assert_not_called()

    async def test_create_excel_file_writes_requested_cells_without_opening_result(self):
        with TemporaryDirectory() as tmp:
            requested = Path(tmp) / "report"
            expected = requested.with_suffix(".xlsx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)

            with mock.patch.object(docgen_module, "_open_file", create=True) as open_file:
                result = await registry.call(
                    "create_excel_file",
                    {
                        "path": str(requested),
                        "headers": ["name", "count"],
                        "rows": [["first", 1], ["second", 2]],
                    },
                )

            self.assertTrue(expected.exists())
            workbook = load_workbook(expected, read_only=True)
            try:
                values = list(workbook.active.values)
            finally:
                workbook.close()
            self.assertEqual([("name", "count"), ("first", 1), ("second", 2)], values)
            open_file.assert_not_called()
            self.assertIn(str(expected), result)

    async def test_excel_formula_syntax_is_preserved_while_plain_chinese_is_normalized(self):
        formulas = ['=IF(A1="支付",1,0)', "='中文表'!A1", '=SUM(Table1[支付])',
                    '=WEBSERVICE("https://example.test/中文?a=1,b=2")']
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / 'formulas.xlsx').resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, '_open_file', create=True) as open_file:
                await registry.call('create_excel_file', {
                    'path': str(path), 'headers': [formulas[0]],
                    'rows': [formulas, ['中文,标点(保持规范)']],
                })
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                values = list(workbook.active.values)
            finally:
                workbook.close()
            self.assertEqual(formulas[0], values[0][0])
            self.assertEqual(tuple(formulas), values[1])
            self.assertEqual('中文，标点（保持规范）', values[2][0])
            open_file.assert_not_called()

    async def test_excel_invalid_arrays_are_rejected_before_writing(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / 'invalid.xlsx').resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            for args in ({'headers': '=WEBSERVICE("x")'}, {'rows': ['=SUM(A1)']},
                         {'rows': '=EXEC("cmd")'}, {'headers': {}}):
                with self.subTest(args=args), mock.patch.object(docgen_module, '_open_file', create=True) as opener:
                    with self.assertRaisesRegex(ValueError, '必须是'):
                        await registry.call('create_excel_file', {'path': str(path), **args})
                    self.assertFalse(path.exists())
                    opener.assert_not_called()

    async def test_excel_protected_directory_is_still_rejected(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / 'protected.xlsx').resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, '_reject_protected',
                                   side_effect=ValueError('受保护目录')), \
                 mock.patch.object(docgen_module, '_open_file', create=True) as opener:
                with self.assertRaisesRegex(ValueError, '受保护目录'):
                    await registry.call('create_excel_file', {'path': str(path), 'rows': [[1]]})
                self.assertFalse(path.exists())
                opener.assert_not_called()

    async def test_create_excel_file_refuses_overwrite_without_opening_or_changing_file(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "existing.xlsx").resolve()
            original = b"existing workbook bytes"
            path.write_bytes(original)
            registry = ToolRegistry()
            docgen_module.register_tools(registry)

            with mock.patch.object(docgen_module, "_open_file", create=True) as open_file:
                with self.assertRaisesRegex(ValueError, "避免覆盖"):
                    await registry.call(
                        "create_excel_file",
                        {"path": str(path), "headers": ["must", "not", "write"]},
                    )

            self.assertEqual(original, path.read_bytes())
            open_file.assert_not_called()


    async def test_create_word_doc_normalizes_halfwidth_punctuation(self):
        """模型写中文常混用半角标点,落盘前必须统一成中文全角。"""
        from docx import Document

        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "方案.docx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                await registry.call("create_word_doc", {
                    "path": str(path),
                    "title": "汽车吊与塔吊对比方案",
                    "paragraphs": ["三栋单体,主体结构施工期3个月(90日历天),完成即退场。"],
                })
            doc = Document(str(path))
            joined = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        self.assertNotIn(",", joined, "半角逗号要换成中文逗号")
        self.assertNotIn("(", joined, "半角左括号要换成中文左括号")
        self.assertIn("，", joined)
        self.assertIn("（", joined)

    async def test_create_word_doc_applies_default_layout(self):
        """默认自动套用 Word 排版规范:25mm 页边距 / 小四宋体 / 固定行距 22 磅 / 章标题小三居中。"""
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "排版.docx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                await registry.call("create_word_doc", {
                    "path": str(path),
                    "header": "火箭基地项目",
                    "blocks": [
                        {"type": "chapter", "text": "一、编制说明"},
                        {"type": "body", "text": "本方案针对三栋单体。"},
                        {"type": "table", "caption": "表1 台班排布",
                         "headers": ["工序", "台班"], "rows": [["吊装", "12"]]},
                    ],
                })
            doc = Document(str(path))
            section = doc.sections[0]
            self.assertAlmostEqual(25.0, section.left_margin.mm, places=1)
            self.assertAlmostEqual(25.0, section.top_margin.mm, places=1)
            self.assertAlmostEqual(0.0, section.gutter.mm, places=1)
            self.assertAlmostEqual(15.0, section.header_distance.mm, places=1)
            self.assertAlmostEqual(17.5, section.footer_distance.mm, places=1)

            chapter = doc.paragraphs[0]
            self.assertEqual("一、编制说明", chapter.text)
            self.assertTrue(chapter.runs[0].font.bold)
            self.assertAlmostEqual(15.0, chapter.runs[0].font.size.pt, places=1)
            self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, chapter.alignment)

            body = doc.paragraphs[1]
            self.assertAlmostEqual(12.0, body.runs[0].font.size.pt, places=1)
            self.assertAlmostEqual(22.0, body.paragraph_format.line_spacing.pt, places=1)
            self.assertGreater(body.paragraph_format.first_line_indent.pt, 0, "正文要有首行缩进")

            self.assertEqual(1, len(doc.tables), "blocks 里的 table 要生成表格")
            self.assertEqual("工序", doc.tables[0].rows[0].cells[0].text)
            self.assertEqual("火箭基地项目", doc.sections[0].header.paragraphs[0].text)

    async def test_create_word_doc_preserves_zero_and_false_in_tables(self):
        from docx import Document

        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "值.docx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                await registry.call("create_word_doc", {
                    "path": str(path),
                    "blocks": [{
                        "type": "table",
                        "headers": ["数字", "布尔值"],
                        "rows": [[0, False]],
                    }],
                })
            doc = Document(str(path))
            values = [cell.text for cell in doc.tables[0].rows[1].cells]
        self.assertEqual(["0", "False"], values)

    async def test_create_word_doc_rejects_malformed_table_rows(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "错误.docx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                with self.assertRaisesRegex(ValueError, "rows.*二维数组"):
                    await registry.call("create_word_doc", {
                        "path": str(path),
                        "blocks": [{
                            "type": "table",
                            "headers": ["列1"],
                            "rows": [None],
                        }],
                    })

    async def test_create_word_doc_sets_table_width_and_refreshable_page_field(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "XML.docx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                await registry.call("create_word_doc", {
                    "path": str(path),
                    "blocks": [{
                        "type": "table",
                        "headers": ["列1", "列2"],
                        "rows": [["a", "b"]],
                    }],
                })
            with ZipFile(path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
                footer_xml = archive.read("word/footer1.xml").decode("utf-8")
                settings_xml = archive.read("word/settings.xml").decode("utf-8")
        self.assertIn('w:tblW w:type="dxa"', document_xml)
        self.assertNotIn('w:tblW w:type="auto" w:w="0"', document_xml)
        self.assertIn('w:fldCharType="separate"', footer_xml)
        self.assertIn("<w:t>1</w:t>", footer_xml)
        self.assertIn("w:updateFields", settings_xml)

    async def test_create_excel_file_normalizes_string_cells_only(self):
        with TemporaryDirectory() as tmp:
            path = (Path(tmp) / "表.xlsx").resolve()
            registry = ToolRegistry()
            docgen_module.register_tools(registry)
            with mock.patch.object(docgen_module, "_open_file", create=True):
                await registry.call("create_excel_file", {
                    "path": str(path),
                    "headers": ["项目,名称", "数量"],
                    "rows": [["钢筋,2200t", 2200]],
                })
            workbook = load_workbook(path, read_only=True)
            try:
                values = list(workbook.active.values)
            finally:
                workbook.close()
        self.assertEqual("项目，名称", values[0][0])
        self.assertEqual("钢筋，2200t", values[1][0])
        self.assertEqual(2200, values[1][1], "数字单元格不能被当文本处理")


if __name__ == "__main__":
    unittest.main()
